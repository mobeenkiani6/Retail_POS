import csv
import io
from datetime import datetime
from flask import jsonify, request, Response
from app.models import db, Sale, Product, Customer, Supplier, Expense, Inventory, User
from app.utils.auth_decorators import token_required, admin_access
from app.routes.admin import admin_bp
from app.routes.admin.helpers import parse_branch_id, period_bounds


def _ensure_sale_invoice(sale):
    """Return display invoice number; persist missing ones so exports stay consistent."""
    if sale.invoice_number:
        return sale.invoice_number, False
    sale.invoice_number = f'INV-{sale.id:06d}'
    db.session.add(sale)
    return sale.invoice_number, True


def _parse_report_bounds():
    period = request.args.get('period', 'month')
    start_date = (request.args.get('start') or request.args.get('from') or '').strip() or None
    end_date = (request.args.get('end') or request.args.get('to') or '').strip() or None
    start, end = period_bounds(period, start_date=start_date, end_date=end_date)
    return period, start, end, start_date, end_date


def _payment_filter():
    raw = (request.args.get('payment_method') or request.args.get('payment') or '').strip()
    if not raw or raw.lower() in ('all', '*'):
        return None
    return raw


def _sales_rows(branch_id, start, end, payment_method=None):
    q = Sale.query.filter(Sale.created_at >= start, Sale.created_at <= end)
    if branch_id:
        q = q.filter(Sale.branch_id == branch_id)
    if payment_method:
        q = q.filter(db.func.lower(Sale.payment_method) == payment_method.lower())
    return q.order_by(Sale.created_at.desc()).all()


def _fmt_dt(value):
    if not value:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M')
    return str(value)[:16].replace('T', ' ')


def _fmt_money(value):
    try:
        return f'{float(value or 0):,.2f}'
    except (TypeError, ValueError):
        return '0.00'


def _header_label(h):
    labels = {
        'invoice': 'Invoice No',
        'payment_method': 'Payment',
        'created_at': 'Date / Time',
        'expense_date': 'Expense Date',
        'last_login_at': 'Last Login',
        'outstanding_balance': 'Outstanding',
        'stock_level': 'Stock',
        'base_price': 'Base Price',
        'cost_price': 'Cost Price',
        'loyalty_points': 'Loyalty Pts',
        'store_credit': 'Store Credit',
        'product_name': 'Product',
        'branch_id': 'Branch',
        'cogs': 'COGS',
        'sku_id': 'SKU ID',
        'supplier_code': 'Code',
    }
    return labels.get(h, h.replace('_', ' ').title())


def _build_report(report_type, branch_id, start, end, payment_method=None):
    rows = []
    headers = []

    if report_type == 'sales':
        headers = ['invoice', 'status', 'payment_method', 'total', 'tax', 'cogs', 'branch_id', 'created_at']
        dirty = False
        for s in _sales_rows(branch_id, start, end, payment_method):
            inv, created = _ensure_sale_invoice(s)
            if created:
                dirty = True
            rows.append([
                inv,
                s.status or '',
                s.payment_method or '',
                round(float(s.total_amount or 0), 2),
                round(float(s.tax_amount or 0), 2),
                round(float(s.cogs_amount or 0), 2),
                s.branch_id or '',
                _fmt_dt(s.created_at),
            ])
        if dirty:
            try:
                db.session.commit()
            except Exception:
                db.session.rollback()
    elif report_type == 'products':
        headers = ['id', 'name', 'sku', 'barcode', 'base_price', 'cost_price', 'status']
        for p in Product.query.filter(Product.archived_at == None).limit(5000).all():
            rows.append([p.id, p.name, p.sku, p.barcode, float(p.base_price or 0), float(p.cost_price or 0), p.status])
    elif report_type == 'customers':
        headers = ['id', 'name', 'phone', 'email', 'loyalty_points', 'store_credit']
        for c in Customer.query.filter(Customer.archived_at == None).limit(5000).all():
            rows.append([c.id, c.name, c.phone, c.email, c.loyalty_points, float(c.store_credit or 0)])
    elif report_type == 'suppliers':
        headers = ['id', 'name', 'code', 'phone', 'city', 'outstanding_balance', 'status']
        for s in Supplier.query.filter(Supplier.archived_at == None).all():
            rows.append([
                s.id, s.name, s.supplier_code, s.phone, s.city,
                float(s.outstanding_balance or 0), s.status,
            ])
    elif report_type == 'expenses':
        headers = ['id', 'title', 'amount', 'payment_method', 'branch_id', 'expense_date']
        q = Expense.query.filter(Expense.expense_date >= start, Expense.expense_date <= end, Expense.archived_at == None)
        if branch_id:
            q = q.filter(Expense.branch_id == branch_id)
        if payment_method:
            q = q.filter(db.func.lower(Expense.payment_method) == payment_method.lower())
        for e in q.order_by(Expense.expense_date.desc()).all():
            rows.append([
                e.id, e.title, float(e.amount or 0), e.payment_method or '', e.branch_id or '',
                _fmt_dt(e.expense_date),
            ])
    elif report_type == 'inventory':
        headers = ['product_id', 'product_name', 'branch_id', 'stock_level', 'sku_id']
        inv = db.session.query(Inventory, Product).join(Product, Inventory.product_id == Product.id).filter(
            Product.archived_at == None
        )
        if branch_id:
            inv = inv.filter(Inventory.branch_id == branch_id)
        for row, p in inv.limit(10000).all():
            rows.append([p.id, p.name, row.branch_id, row.stock_level, row.sku_id])
    elif report_type == 'employees':
        headers = ['id', 'username', 'role', 'branch_id', 'last_login_at']
        for u in User.query.filter(User.archived_at == None).all():
            rows.append([u.id, u.username, u.role, u.branch_id, _fmt_dt(u.last_login_at)])
    else:
        return None, None
    return headers, rows


def _period_label(period, start, end, start_date, end_date):
    if period == 'custom':
        a = start_date or (start.strftime('%Y-%m-%d') if start else '')
        b = end_date or (end.strftime('%Y-%m-%d') if end else '')
        return f'{a} → {b}'
    labels = {
        'today': 'Today',
        'week': 'This week',
        'month': 'This month',
        'quarter': 'This quarter',
        'year': 'This year',
    }
    return labels.get(period, period)


def _build_pdf(report_type, headers, rows, meta):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    out = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(
        out,
        pagesize=page,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=2,
        leading=22,
    )
    brand_style = ParagraphStyle(
        'Brand',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        textColor=colors.HexColor('#0d9488'),
        spaceAfter=6,
        tracking=1,
    )
    meta_style = ParagraphStyle(
        'Meta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        textColor=colors.HexColor('#64748b'),
        leading=12,
    )
    cell_style = ParagraphStyle(
        'Cell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        textColor=colors.HexColor('#1e293b'),
        leading=10,
    )
    head_cell_style = ParagraphStyle(
        'HeadCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        textColor=colors.white,
        leading=10,
    )

    money_cols = {i for i, h in enumerate(headers) if h in (
        'total', 'tax', 'cogs', 'amount', 'base_price', 'cost_price',
        'store_credit', 'outstanding_balance',
    )}

    display_headers = [Paragraph(_header_label(h), head_cell_style) for h in headers]
    table_data = [display_headers]

    for r in rows:
        cells = []
        for i, val in enumerate(r):
            if i in money_cols:
                text = _fmt_money(val)
            else:
                text = '' if val is None else str(val)
            # Keep cells compact for layout
            if len(text) > 48:
                text = text[:45] + '…'
            cells.append(Paragraph(text.replace('&', '&amp;').replace('<', '&lt;'), cell_style))
        table_data.append(cells)

    available = page[0] - 28 * mm
    n = max(len(headers), 1)
    # Prefer wider first columns (invoice / name)
    weights = []
    for h in headers:
        if h in ('invoice', 'name', 'title', 'product_name', 'email', 'address'):
            weights.append(1.6)
        elif h in ('created_at', 'expense_date', 'last_login_at', 'branch_id'):
            weights.append(1.3)
        elif h in ('id', 'status', 'sku', 'code'):
            weights.append(0.8)
        else:
            weights.append(1.0)
    total_w = sum(weights)
    col_widths = [available * (w / total_w) for w in weights]

    accent = colors.HexColor('#0d9488')
    header_bg = colors.HexColor('#0f766e')
    stripe = colors.HexColor('#f0fdfa')
    grid = colors.HexColor('#e2e8f0')
    text_muted = colors.HexColor('#64748b')

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), header_bg),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('ALIGN', (0, 0), (-1, 0), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.4, grid),
        ('BOX', (0, 0), (-1, -1), 1, accent),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, stripe]),
    ]
    for i in money_cols:
        style_cmds.append(('ALIGN', (i, 1), (i, -1), 'RIGHT'))
    table.setStyle(TableStyle(style_cmds))

    report_title = f'{report_type.replace("_", " ").title()} Report'
    period_txt = meta.get('period_label', '')
    payment_txt = meta.get('payment_label') or 'All payments'
    branch_txt = meta.get('branch_label') or 'All branches'
    generated = datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')
    row_count = len(rows)

    summary_bits = [f'<b>{row_count}</b> rows', f'Period: <b>{period_txt}</b>', f'Branch: {branch_txt}']
    if meta.get('payment_label'):
        summary_bits.append(f'Payment: <b>{payment_txt}</b>')
    if meta.get('sales_total') is not None:
        summary_bits.append(f'Total sales: <b>{_fmt_money(meta["sales_total"])}</b>')

    story = [
        Paragraph('NYCTO RETAIL', brand_style),
        Paragraph(report_title, title_style),
        Paragraph(' · '.join(summary_bits), meta_style),
        Paragraph(f'Generated {generated}', meta_style),
        Spacer(1, 8),
    ]

    if rows:
        story.append(table)
    else:
        empty = ParagraphStyle('Empty', parent=meta_style, alignment=TA_CENTER, fontSize=11, textColor=text_muted)
        story.append(Spacer(1, 40))
        story.append(Paragraph('No rows match the selected filters.', empty))

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#e2e8f0'))
        canvas.setLineWidth(0.6)
        canvas.line(14 * mm, 11 * mm, page[0] - 14 * mm, 11 * mm)
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(colors.HexColor('#94a3b8'))
        canvas.drawString(14 * mm, 6 * mm, 'Nycto Retail · Confidential')
        canvas.drawRightString(page[0] - 14 * mm, 6 * mm, f'Page {doc_.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    out.seek(0)
    return out


@admin_bp.route('/reports/summary', methods=['GET'])
@token_required
@admin_access
def reports_summary(current_user):
    period, start, end, start_date, end_date = _parse_report_bounds()
    branch_id = parse_branch_id()
    payment_method = _payment_filter()
    sales = _sales_rows(branch_id, start, end, payment_method)
    revenue = sum(float(s.total_amount or 0) for s in sales if s.status in ('completed', 'partially_returned'))
    cogs = sum(float(s.cogs_amount or 0) for s in sales if s.status in ('completed', 'partially_returned'))
    refunded = sum(float(s.total_amount or 0) for s in sales if s.status == 'refunded')
    exp_q = Expense.query.filter(Expense.expense_date >= start, Expense.expense_date <= end, Expense.archived_at == None)
    if branch_id:
        exp_q = exp_q.filter(Expense.branch_id == branch_id)
    if payment_method:
        exp_q = exp_q.filter(db.func.lower(Expense.payment_method) == payment_method.lower())
    expenses = sum(float(e.amount or 0) for e in exp_q.all())
    return jsonify({
        'period': period,
        'start': start.isoformat(),
        'end': end.isoformat(),
        'payment_method': payment_method or 'all',
        'revenue': round(revenue, 2),
        'cost_of_goods_sold': round(cogs, 2),
        'cogs': round(cogs, 2),
        'gross_profit': round(revenue - cogs, 2),
        'expenses': round(expenses, 2),
        'net_profit': round(revenue - cogs - expenses, 2),
        'orders': len(sales),
        'refunded_amount': round(refunded, 2),
    }), 200


@admin_bp.route('/reports/preview', methods=['GET'])
@token_required
@admin_access
def reports_preview(current_user):
    report_type = request.args.get('type', 'sales')
    period, start, end, start_date, end_date = _parse_report_bounds()
    branch_id = parse_branch_id()
    payment_method = _payment_filter()
    headers, rows = _build_report(report_type, branch_id, start, end, payment_method)
    if headers is None:
        return jsonify({'message': f'Unknown report type: {report_type}'}), 400
    limit = min(int(request.args.get('limit', 50)), 200)
    return jsonify({
        'type': report_type,
        'period': period,
        'payment_method': payment_method or 'all',
        'headers': headers,
        'rows': rows[:limit],
        'total_rows': len(rows),
        'preview_rows': min(limit, len(rows)),
        'start': start.isoformat(),
        'end': end.isoformat(),
    }), 200


@admin_bp.route('/reports/export', methods=['GET'])
@token_required
@admin_access
def reports_export(current_user):
    report_type = request.args.get('type', 'sales')
    fmt = request.args.get('format', 'csv')
    period, start, end, start_date, end_date = _parse_report_bounds()
    branch_id = parse_branch_id()
    payment_method = _payment_filter()

    headers, rows = _build_report(report_type, branch_id, start, end, payment_method)
    if headers is None:
        return jsonify({'message': f'Unknown report type: {report_type}'}), 400

    period_slug = period
    if period == 'custom' and start_date and end_date:
        period_slug = f'{start_date}_{end_date}'
    pay_slug = f'_{payment_method}' if payment_method else ''
    filename_base = f'{report_type}_{period_slug}{pay_slug}'

    if fmt == 'json':
        return jsonify({'headers': headers, 'rows': rows, 'total_rows': len(rows)}), 200

    if fmt == 'csv':
        # Human-friendly header row
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([_header_label(h) for h in headers])
        writer.writerows(rows)
        return Response(
            buf.getvalue(),
            mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename={filename_base}.csv'},
        )

    if fmt in ('xlsx', 'excel'):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            wb = Workbook()
            ws = wb.active
            ws.title = report_type[:31]
            pretty = [_header_label(h) for h in headers]
            ws.append(pretty)
            for r in rows:
                ws.append(list(r))
            header_fill = PatternFill('solid', fgColor='0F766E')
            header_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
            thin = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0'),
            )
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='left', vertical='center')
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
                for cell in row:
                    cell.border = thin
                    cell.alignment = Alignment(vertical='center')
            for col in ws.columns:
                max_len = 0
                letter = col[0].column_letter
                for cell in col:
                    max_len = max(max_len, len(str(cell.value or '')))
                ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 36)
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = 'A2'
            out = io.BytesIO()
            wb.save(out)
            out.seek(0)
            return Response(
                out.getvalue(),
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                headers={'Content-Disposition': f'attachment; filename={filename_base}.xlsx'},
            )
        except ImportError:
            return jsonify({'message': 'openpyxl not installed; use format=csv'}), 501

    if fmt == 'pdf':
        try:
            sales_total = None
            if report_type == 'sales' and 'total' in headers:
                ti = headers.index('total')
                sales_total = sum(float(r[ti] or 0) for r in rows)
            meta = {
                'period_label': _period_label(period, start, end, start_date, end_date),
                'payment_label': payment_method,
                'branch_label': branch_id or 'All branches',
                'sales_total': sales_total,
            }
            out = _build_pdf(report_type, headers, rows, meta)
            return Response(
                out.getvalue(),
                mimetype='application/pdf',
                headers={'Content-Disposition': f'attachment; filename={filename_base}.pdf'},
            )
        except ImportError:
            return jsonify({'message': 'reportlab not installed; use format=csv'}), 501
        except Exception as e:
            return jsonify({'message': f'PDF export failed: {e}'}), 500

    return jsonify({'message': 'Unsupported format'}), 400
